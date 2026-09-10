import tkinter as tk
from ui.style import *
from tkinter import ttk, messagebox, filedialog
import sqlite3
import pandas as pd
import os  # Para abrir el archivo automáticamente
from datetime import datetime
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from data.inventario_db import InventarioDB
import winsound
import textwrap

class VentanaOrdenes:
    def __init__(self, root, main_window):
        self.root = root
        self.main_window = main_window
        self.db = InventarioDB()
        
        self.root.title("GENERADOR DE ÓRDENES DE SALIDA")
        center_window(self.root, *MODULE_WINDOW_SIZE)
        self.root.configure(bg="#f4f6f7")
        self.root.protocol("WM_DELETE_WINDOW", self.regresar)

        # --- ENCABEZADO ---
        header = tk.Frame(root, bg="#2980b9", pady=15)
        header.pack(fill="x")
        
        tk.Button(header, text="⬅ VOLVER", command=self.regresar, bg="#1a5276", fg="white", 
                  font=FONT_LABEL_BOLD).pack(side="left", padx=20)
        
        tk.Label(header, text="GESTIÓN DE DOCUMENTOS DE SALIDA", bg="#2980b9", 
                 fg="white", font=FONT_HEADER).pack(side="left", padx=20)

        # --- FILTROS ---
        frame_filtros = tk.LabelFrame(root, text=" 1. CONFIGURACIÓN DE LA ORDEN ", padx=20, pady=15, bg="white")
        frame_filtros.pack(fill="x", padx=30, pady=20)

        tk.Label(frame_filtros, text="MUNICIPIO:", bg="white").grid(row=0, column=0, sticky="w")
        self.combo_mun = ttk.Combobox(frame_filtros, state="readonly", width=35)
        self.combo_mun.grid(row=0, column=1, padx=10)
        self.combo_mun['values'] = self.db.obtener_municipios()

        tk.Label(frame_filtros, text="PERIODO:", bg="white").grid(row=0, column=2, sticky="w", padx=10)
        self.combo_per = ttk.Combobox(frame_filtros, values=("PRIMAVERA", "VERANO", "OTOÑO"), state="readonly", width=15)
        self.combo_per.grid(row=0, column=3, padx=10)

        tk.Label(frame_filtros, text="AÑO:", bg="white").grid(row=0, column=4, sticky="w", padx=10)
        self.combo_anio = ttk.Combobox(frame_filtros, values=[str(a) for a in range(2024, 2035)], state="readonly", width=10)
        self.combo_anio.set(str(datetime.now().year))
        self.combo_anio.grid(row=0, column=5, padx=10)
    
        # --- TABLA ---
        frame_tabla = tk.LabelFrame(root, text=" 2. VISTA PREVIA DE PRODUCTOS ", bg="white")
        frame_tabla.pack(fill="both", expand=True, padx=30, pady=10)

        columnas = ("SKU", "TIPO", "MEDICAMENTO", "CANTIDAD")
        self.tabla = ttk.Treeview(frame_tabla, columns=columnas, show="headings")
        for col in columnas:
            self.tabla.heading(col, text=col)
            self.tabla.column(col, anchor="center")
        
        
        # Ajustamos anchos
        self.tabla.column("TIPO", width=120, anchor="center")
        self.tabla.column("MEDICAMENTO", width=380, anchor="w")
        self.tabla.pack(fill="both", expand=True, padx=10, pady=10)

        self.tabla.tag_configure('inactivo', foreground='#c0392b')

        # --- BOTONES ---
        btn_frame = tk.Frame(root, bg="#f4f6f7")
        btn_frame.pack(pady=20)
        tk.Button(frame_filtros, text="🔍 CARGAR DATOS", command=self.cargar_vista_previa, 
                  bg="#3498db", fg="white", font=FONT_LABEL_BOLD).grid(row=0, column=6, padx=20)
        tk.Button(btn_frame, text="📄 GENERAR ORDEN PDF", command=self.generar_pdf, 
                  bg="#e74c3c", fg="white", font=FONT_BUTTON_LARGE, width=25, height=2).pack(side="left", padx=15)

        tk.Button(btn_frame, text="Excel GENERAR EXCEL", command=self.generar_excel, 
                  bg="#27ae60", fg="white", font=FONT_BUTTON_LARGE, width=25, height=2).pack(side="left", padx=15)

        tk.Button(btn_frame, text="🗂️ IMPRIMIR TODO LOTE", command=self.generar_todos_pdf, 
                  bg="#34495e", fg="white", font=FONT_BUTTON_LARGE, width=22, height=2).pack(side="left", padx=10)
    
        tk.Button(btn_frame, text="📊 EXCEL TODO LOTE", command=self.generar_todos_excel, 
          bg="#1d7344", fg="white", font=FONT_BUTTON_LARGE, width=22, height=2).pack(side="left", padx=10)
    def cargar_vista_previa(self):
        mun = self.combo_mun.get()
        per = self.combo_per.get()
        anio = self.combo_anio.get()

        if not mun or not per:
            messagebox.showwarning("Atención", "Seleccione Municipio y Periodo"); return

        for i in self.tabla.get_children(): self.tabla.delete(i)

        conn = sqlite3.connect(self.db.db_name)
        query = """
            SELECT h.sku, m.nombre || ' ' || m.gramaje || ' - ' || m.presentacion, h.cantidad, m.estatus
            FROM historial_retornos h
            JOIN medicamentos m ON h.sku = m.sku
            WHERE h.municipio = ? AND h.periodo = ? AND h.anio = ? AND h.tipo_movimiento = 'SALIDA MUNICIPIO'
        """
        cursor = conn.cursor()
        cursor.execute(query, (mun, per, anio))
        filas = cursor.fetchall()
        
        self.datos_orden = []
        for fila in filas:
            sku, nombre, cant, estatus = fila
            
            if estatus == 'INACTIVO':
                # CORRECCIÓN: Usamos [!] en lugar del emoji ⚠️ para que el PDF no falle
                nombre_con_aviso = f"[!] {nombre} (INACTIVO)"
                self.datos_orden.append((sku, nombre_con_aviso, cant))
                self.tabla.insert("", "end", values=(sku, nombre_con_aviso, cant), tags=('inactivo',))
            else:
                self.datos_orden.append((sku, nombre, cant))
                self.tabla.insert("", "end", values=(sku, nombre, cant))
                
        conn.close()
    def cargar_vista_previa(self):
        mun = self.combo_mun.get()
        per = self.combo_per.get()
        anio = self.combo_anio.get()

        if not mun or not per:
            messagebox.showwarning("Atención", "Seleccione Municipio y Periodo"); return

        for i in self.tabla.get_children(): self.tabla.delete(i)

        conn = sqlite3.connect(self.db.db_name)
        # Añadido m.tipo y ORDER BY
        query = """
            SELECT h.sku, m.tipo, m.nombre || ' ' || m.gramaje || ' - ' || m.presentacion, h.cantidad, m.estatus
            FROM historial_retornos h
            JOIN medicamentos m ON h.sku = m.sku
            WHERE h.municipio = ? AND h.periodo = ? AND h.anio = ? AND h.tipo_movimiento = 'SALIDA MUNICIPIO'
            ORDER BY m.tipo ASC, m.nombre ASC
        """
        cursor = conn.cursor()
        cursor.execute(query, (mun, per, anio))
        filas = cursor.fetchall()
        
        self.datos_orden = []
        for fila in filas:
            sku, tipo, nombre, cant, estatus = fila # Ahora son 5 variables
            
            if estatus == 'INACTIVO':
                nombre_con_aviso = f"[!] {nombre} (INACTIVO)"
                self.datos_orden.append((sku, tipo, nombre_con_aviso, cant))
                self.tabla.insert("", "end", values=(sku, tipo, nombre_con_aviso, cant), tags=('inactivo',))
            else:
                self.datos_orden.append((sku, tipo, nombre, cant))
                self.tabla.insert("", "end", values=(sku, tipo, nombre, cant))
                
        conn.close()

    def generar_pdf(self):
        if not hasattr(self, 'datos_orden') or not self.datos_orden:
            messagebox.showerror("Error", "No hay datos cargados."); return

        # NUEVO: Generar un nombre por defecto basado en municipio y fecha/hora
        mun_limpio = self.combo_mun.get().split(" - ")[0] # Toma solo el nombre del municipio
        fecha_str = datetime.now().strftime("%Y%m%d_%H%M")
        nombre_default = f"Orden_{mun_limpio}_{fecha_str}.pdf"

        filepath = filedialog.asksaveasfilename(
            defaultextension=".pdf", 
            initialfile=nombre_default, # Se asigna el nombre sugerido
            filetypes=[("PDF files", "*.pdf")]
        )
        if not filepath: return

        c = canvas.Canvas(filepath, pagesize=letter)
        # Encabezado del PDF
        c.setFont("Helvetica-Bold", 16)
        c.drawString(50, 750, "PLANASSZE IPN - ORDEN DE SALIDA A MUNICIPIO")
        c.setFont("Helvetica", 10)
        c.drawString(50, 735, f"DESTINO: {self.combo_mun.get()}")
        c.drawString(50, 720, f"PERIODO: {self.combo_per.get()} {self.combo_anio.get()}")
        c.drawString(430, 735, f"Generado el: {datetime.now().strftime('%d/%m/%Y %H:%M')}")
        
        c.line(50, 710, 550, 710)
        
        # Busca este bloque de código dentro de generar_pdf Y dibujar_pdf_individual
        # y reemplázalo por este:
        
        # Encabezado Tabla
        y = 680
        c.setFont("Helvetica-Bold", 10)
        c.drawString(50, y, "SKU")
        c.drawString(120, y, "TIPO")          # Nueva columna
        c.drawString(220, y, "MEDICAMENTO")   # Recorrido a la derecha
        c.drawString(500, y, "CANT")
        # Si estás en dibujar_pdf_individual, recuerda que tienes una línea extra aquí:
        # c.line(50, y-5, 550, y-5)
        
        y -= 20
        c.setFont("Helvetica", 9)
        
        # Ojo: usa self.datos_orden en generar_pdf y 'datos' en dibujar_pdf_individual
        for sku, tipo, nombre, cant in self.datos_orden: 
            import textwrap 
            lineas_nombre = textwrap.wrap(str(nombre), width=45) # Reducimos width

            c.drawString(50, y, str(sku))
            c.drawString(120, y, str(tipo))      # Imprimimos el tipo
            c.drawString(500, y, str(cant))

            y_temp = y
            for linea in lineas_nombre:
                c.drawString(220, y_temp, linea) # Imprimimos nombre en x=220
                y_temp -= 12 

            y = y_temp - 8
            
            if y < 120:
                c.showPage()
                y = 750
                c.setFont("Helvetica", 9) # Restaurar fuente
        
        # --- ZONA DE FIRMAS ---
        y -= 40
        if y < 80: # Si ya no cabe la firma en esta hoja, brincamos a la siguiente
            c.showPage()
            y = 700

        # LADO IZQUIERDO: ENTREGAS
        c.line(70, y, 220, y)  # Línea de firma
        c.drawCentredString(145, y - 15, "ENTREGA ALMACÉN") # Texto centrado bajo la línea

        # LADO DERECHO: RECIBE
        c.line(350, y, 500, y) # Línea de firma
        c.drawCentredString(425, y - 15, "RECIBE:") 
        c.setFont("Helvetica", 8) 
        c.drawCentredString(425, y - 28, "NOMBRE Y FIRMA") # Segunda línea manual
        
        c.save()
        messagebox.showinfo("Éxito", "PDF Generado correctamente.")
        os.startfile(filepath)

    def generar_excel(self):
        if not hasattr(self, 'datos_orden') or not self.datos_orden:
            messagebox.showerror("Error", "No hay datos cargados."); return

        # Configuración de nombre de archivo
        mun_limpio = self.combo_mun.get().split(" - ")[0]
        fecha_actual = datetime.now().strftime('%d/%m/%Y %H:%M')
        nombre_default = f"Orden_{mun_limpio}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"

        filepath = filedialog.asksaveasfilename(
            defaultextension=".xlsx", 
            initialfile=nombre_default,
            filetypes=[("Excel files", "*.xlsx")]
        )
        if not filepath: return

        try:
            # 1. Crear el DataFrame con las 4 columnas actualizadas
            df = pd.DataFrame(self.datos_orden, columns=["SKU", "TIPO", "MEDICAMENTO", "CANTIDAD"])

            # 2. Iniciar el Escritor (dejamos 5 filas libres para el encabezado)
            writer = pd.ExcelWriter(filepath, engine='openpyxl')
            df.to_excel(writer, index=False, sheet_name='Orden de Salida', startrow=5)

            workbook  = writer.book
            worksheet = writer.sheets['Orden de Salida']

            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

            # --- ESTILOS ---
            title_font = Font(name="Arial", size=14, bold=True, color="2980B9")
            header_fill = PatternFill(start_color="2980B9", end_color="2980B9", fill_type="solid")
            header_font = Font(color="FFFFFF", bold=True)
            border_thin = Border(left=Side(style='thin'), right=Side(style='thin'), 
                                 top=Side(style='thin'), bottom=Side(style='thin'))
            center_align = Alignment(horizontal="center")

            # --- 3. ENCABEZADO (TÍTULOS SUPERIORES) ---
            worksheet["A1"] = "PLANASSZE IPN - ORDEN DE SALIDA A MUNICIPIO"
            worksheet["A1"].font = title_font
            
            worksheet["A2"] = f"DESTINO: {self.combo_mun.get()}"
            worksheet["A3"] = f"PERIODO: {self.combo_per.get()} {self.combo_anio.get()}"
            worksheet["A4"] = f"FECHA DE GENERACIÓN: {fecha_actual}"
            
            for i in range(2, 5):
                worksheet[f"A{i}"].font = Font(bold=True)

            # --- 4. DISEÑO DE LA TABLA DE DATOS ---
            # La tabla empieza en la fila 6 (startrow=5 + 1)
            header_row = 6
            for cell in worksheet[header_row]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = center_align
                cell.border = border_thin

            # Ajuste de columnas
            worksheet.column_dimensions['A'].width = 15 # Columna SKU
            worksheet.column_dimensions['B'].width = 18 # Columna TIPO
            worksheet.column_dimensions['C'].width = 50 # Columna MEDICAMENTO
            worksheet.column_dimensions['D'].width = 15 # Columna CANTIDAD

            # Aplicar bordes y alineación a los datos
            last_row = header_row + len(df)
            for row in worksheet.iter_rows(min_row=header_row + 1, max_row=last_row):
                for cell in row:
                    cell.border = border_thin
                    # Centrar todo excepto el nombre del medicamento (columna 3 / C)
                    if cell.column != 3: 
                        cell.alignment = center_align

            # --- 5. CUADRO DE FIRMAS (AL FINAL DE LA TABLA) ---
            y_firmas = last_row + 4
            
            # Líneas de firma (usando bordes inferiores). Se movió la firma derecha a la columna 4 (D)
            worksheet.cell(row=y_firmas, column=1).border = Border(bottom=Side(style='medium'))
            worksheet.cell(row=y_firmas, column=4).border = Border(bottom=Side(style='medium'))
            
            # Etiquetas de firma
            worksheet.cell(row=y_firmas + 1, column=1, value="ENTREGA ALMACÉN").alignment = center_align
            worksheet.cell(row=y_firmas + 1, column=1).font = Font(bold=True)
            
            worksheet.cell(row=y_firmas + 1, column=4, value="RECIBE (NOMBRE Y FIRMA)").alignment = center_align
            worksheet.cell(row=y_firmas + 1, column=4).font = Font(bold=True)

            # 6. Guardar y abrir
            writer.close()
            messagebox.showinfo("Éxito", "Excel generado con encabezados y firmas.")
            os.startfile(filepath)

        except Exception as e:
            messagebox.showerror("Error", f"No se pudo crear el Excel: {e}")    
    def generar_todos_pdf(self):
        per = self.combo_per.get()
        anio = self.combo_anio.get()

        # 1. VALIDACIÓN PREVIA (Antes de pedir carpeta)
        if not per:
            winsound.MessageBeep(winsound.MB_ICONHAND)
            messagebox.showwarning("Atención", "Seleccione el Periodo para validar el lote.")
            return

        try:
            conn = sqlite3.connect(self.db.db_name)
            cursor = conn.cursor()
            
            # --- CORRECCIÓN: 1. Primero buscamos qué municipios tienen datos ---
            cursor.execute("""
                SELECT DISTINCT municipio 
                FROM historial_retornos 
                WHERE periodo = ? AND anio = ? AND tipo_movimiento = 'SALIDA MUNICIPIO'
            """, (per, anio))
            
            municipios = [m[0] for m in cursor.fetchall()]
            
            # 2. VERIFICACIÓN DE DATOS
            if not municipios:
                winsound.MessageBeep(winsound.MB_ICONHAND)
                messagebox.showinfo("Sin Datos", f"No existen registros para {per} {anio}.")
                conn.close()
                return

            # 3. VISTA PREVIA EN LA TABLA
            for i in self.tabla.get_children(): self.tabla.delete(i)
            self.tabla.heading("SKU", text="MUNICIPIO")
            self.tabla.heading("MEDICAMENTO", text="ESTADO")
            self.tabla.heading("CANTIDAD", text="ITEMS")

            for mun in municipios:
                cursor.execute("SELECT COUNT(*) FROM historial_retornos WHERE municipio=? AND periodo=? AND anio=?", (mun, per, anio))
                total = cursor.fetchone()[0]
                self.tabla.insert("", "end", values=(mun, "LISTO PARA PDF", f"{total} productos"))

            # 4. CONFIRMACIÓN Y CARPETA
            if not messagebox.askyesno("Confirmar Lote", f"Se encontraron {len(municipios)} municipios.\n¿Desea generar los archivos?"):
                conn.close()
                self.reset_tabla_original()
                return

            folder_selected = filedialog.askdirectory(title="Seleccione carpeta de destino")
            if not folder_selected: 
                conn.close()
                return

            # --- CORRECCIÓN: 5. GENERACIÓN MASIVA (Aquí va la consulta con la presentación) ---
            # Reemplaza la consulta SQL dentro de tu ciclo for en ambas funciones
            for mun in municipios:
                cursor.execute("""
                    SELECT h.sku, m.tipo, m.nombre || ' ' || m.gramaje || ' - ' || m.presentacion, h.cantidad
                    FROM historial_retornos h
                    JOIN medicamentos m ON h.sku = m.sku
                    WHERE h.municipio = ? AND h.periodo = ? AND h.anio = ? AND h.tipo_movimiento = 'SALIDA MUNICIPIO'
                    ORDER BY m.tipo ASC, m.nombre ASC
                """, (mun, per, anio))
                # ... resto del código ...
                datos_mun = cursor.fetchall()
                
                # Nombre de archivo seguro (quita caracteres prohibidos en Windows)
                mun_safe = "".join([c for c in mun if c.isalnum() or c in (' ', '_')]).strip()
                filename = os.path.join(folder_selected, f"Orden_{mun_safe}_{per}_{anio}.pdf")
                
                # AQUÍ SE LLAMA A LA FUNCIÓN DE DIBUJO
                self.dibujar_pdf_individual(filename, mun, per, anio, datos_mun)

            conn.close()
            winsound.Beep(1500, 100); winsound.Beep(2000, 150)
            messagebox.showinfo("Éxito", "Lote generado correctamente.")
            os.startfile(folder_selected)
            self.reset_tabla_original()

        except Exception as e:
            messagebox.showerror("Error", f"Fallo en proceso por lotes: {e}")
    # --- FUNCIÓN DE DIBUJO (LA QUE FALTABA) ---
    # --- FUNCIÓN DE DIBUJO (CON SALTO DE LÍNEA AUTOMÁTICO) ---
    def dibujar_pdf_individual(self, filepath, mun, per, anio, datos):
        c = canvas.Canvas(filepath, pagesize=letter)
        ahora = datetime.now().strftime('%d/%m/%Y %H:%M') 
        
        # Encabezado
        c.setFont("Helvetica-Bold", 16)
        c.drawString(50, 750, "PLANASSZE IPN - ORDEN DE SALIDA")
        c.setFont("Helvetica", 10)
        c.drawString(50, 735, f"DESTINO: {mun}")
        c.drawString(50, 720, f"PERIODO: {per} {anio}")
        c.drawString(430, 735, f"Generado: {ahora}")
        c.line(50, 710, 550, 710)

        # Busca este bloque de código dentro de generar_pdf Y dibujar_pdf_individual
        # y reemplázalo por este:
        
        # Encabezado Tabla
        y = 680
        c.setFont("Helvetica-Bold", 10)
        c.drawString(50, y, "SKU")
        c.drawString(120, y, "TIPO")          # Nueva columna
        c.drawString(220, y, "MEDICAMENTO")   # Recorrido a la derecha
        c.drawString(500, y, "CANT")
        # Si estás en dibujar_pdf_individual, recuerda que tienes una línea extra aquí:
        # c.line(50, y-5, 550, y-5)
        
        y -= 20
        c.setFont("Helvetica", 9)
        
        # Ojo: usa self.datos_orden en generar_pdf y 'datos' en dibujar_pdf_individual
        for sku, tipo, nombre, cant in datos:
            import textwrap 
            lineas_nombre = textwrap.wrap(str(nombre), width=45) # Reducimos width

            c.drawString(50, y, str(sku))
            c.drawString(120, y, str(tipo))      # Imprimimos el tipo
            c.drawString(500, y, str(cant))

            y_temp = y
            for linea in lineas_nombre:
                c.drawString(220, y_temp, linea) # Imprimimos nombre en x=220
                y_temp -= 12 

            y = y_temp - 8
            
            if y < 120:
                c.showPage()
                y = 750
                c.setFont("Helvetica", 9)
        # Firmas al final
        y -= 40
        if y < 80: 
            c.showPage()
            y = 700
            
        c.line(70, y, 220, y)
        c.drawCentredString(145, y - 15, "ENTREGA ALMACÉN")
        c.line(350, y, 500, y)
        c.drawCentredString(425, y - 15, "RECIBE:")
        c.drawCentredString(425, y - 28, "NOMBRE Y FIRMA")
        
        c.save()

    def generar_todos_excel(self):
        """Valida, muestra vista previa y genera Excels individuales por municipio en lote"""
        per = self.combo_per.get()
        anio = self.combo_anio.get()

        if not per:
            winsound.MessageBeep(winsound.MB_ICONHAND)
            messagebox.showwarning("Atención", "Seleccione el Periodo para validar el lote."); return

        try:
            conn = sqlite3.connect(self.db.db_name)
            cursor = conn.cursor()
            
            # --- CORRECCIÓN: 1. Buscar municipios con movimientos (Restaurado) ---
            cursor.execute("""
                SELECT DISTINCT municipio FROM historial_retornos 
                WHERE periodo = ? AND anio = ? AND tipo_movimiento = 'SALIDA MUNICIPIO'
            """, (per, anio))
            municipios = [m[0] for m in cursor.fetchall()]
            
            if not municipios:
                winsound.MessageBeep(winsound.MB_ICONHAND)
                messagebox.showinfo("Sin Datos", "No hay registros para este periodo.")
                conn.close(); return

            # 2. Vista previa en tabla
            self.mostrar_vista_previa_lote(cursor, municipios, per, anio)

            # 3. Confirmación y Carpeta
            if not messagebox.askyesno("Confirmar Lote Excel", f"Se generarán {len(municipios)} libros de Excel.\n¿Continuar?"):
                conn.close(); self.reset_tabla_original(); return

            folder = filedialog.askdirectory(title="Carpeta para guardar los Excels")
            if not folder: conn.close(); return

            # --- CORRECCIÓN: 4. Generación Masiva (Con la presentación incluida) ---
            # Reemplaza la consulta SQL dentro de tu ciclo for en ambas funciones
            for mun in municipios:
                cursor.execute("""
                    SELECT h.sku, m.tipo, m.nombre || ' ' || m.gramaje || ' - ' || m.presentacion, h.cantidad
                    FROM historial_retornos h
                    JOIN medicamentos m ON h.sku = m.sku
                    WHERE h.municipio = ? AND h.periodo = ? AND h.anio = ? AND h.tipo_movimiento = 'SALIDA MUNICIPIO'
                    ORDER BY m.tipo ASC, m.nombre ASC
                """, (mun, per, anio))
                # ... resto del código ...
                datos = cursor.fetchall()
                
                mun_safe = "".join([c for c in mun if c.isalnum() or c in (' ', '_')]).strip()
                filepath = os.path.join(folder, f"Orden_{mun_safe}_{per}_{anio}.xlsx")
                
                # Función auxiliar de escritura
                self.escribir_excel_individual(filepath, mun, per, anio, datos)

            conn.close()
            winsound.Beep(1500, 100); winsound.Beep(2000, 150)
            messagebox.showinfo("Éxito", "Lote de Excel generado correctamente.")
            os.startfile(folder)
            self.reset_tabla_original()

        except Exception as e:
            messagebox.showerror("Error", f"Fallo en lote Excel: {e}")


    def escribir_excel_individual(self, filepath, mun, per, anio, datos):
        """Crea un archivo Excel con formato profesional para un municipio"""
        import pandas as pd
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        ahora = datetime.now().strftime('%d/%m/%Y %H:%M') 
        
        # 1. Se agrega la columna TIPO al DataFrame
        df = pd.DataFrame(datos, columns=["SKU", "TIPO", "MEDICAMENTO", "CANTIDAD"])
        writer = pd.ExcelWriter(filepath, engine='openpyxl')
        df.to_excel(writer, index=False, sheet_name='Orden', startrow=5)

        ws = writer.sheets['Orden']
        
        # Títulos
        ws["A1"] = "PLANASSZE IPN - ORDEN DE SALIDA"
        ws["A1"].font = Font(size=14, bold=True, color="2980B9")
        ws["A2"] = f"DESTINO: {mun}"
        ws["A3"] = f"PERIODO: {per} {anio}"
        ws["A4"] = f"FECHA Y HORA DE EMISIÓN: {ahora}" 
        ws["A4"].font = Font(bold=True)
        
        # Estilos Tabla
        header_fill = PatternFill(start_color="2980B9", end_color="2980B9", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        center_align = Alignment(horizontal="center")

        for cell in ws[6]: # Fila de encabezados
            cell.fill, cell.font, cell.border, cell.alignment = header_fill, header_font, border, center_align

        # 2. Ajuste para las 4 columnas
        ws.column_dimensions['A'].width = 15 # SKU
        ws.column_dimensions['B'].width = 18 # TIPO
        ws.column_dimensions['C'].width = 50 # MEDICAMENTO
        ws.column_dimensions['D'].width = 15 # CANTIDAD

        last_row = 6 + len(df)
        
        # 3. Aplicar bordes y centrado a las celdas de datos
        for row in ws.iter_rows(min_row=7, max_row=last_row):
            for cell in row:
                cell.border = border
                if cell.column != 3: # Centrar todo excepto el nombre del medicamento
                    cell.alignment = center_align

        # 4. Firmas (Se mueve la firma derecha a la columna 4)
        y_f = last_row + 4
        ws.cell(row=y_f, column=1).border = Border(bottom=Side(style='medium'))
        ws.cell(row=y_f, column=4).border = Border(bottom=Side(style='medium')) # <-- Cambio a col 4
        
        ws.cell(row=y_f+1, column=1, value="ENTREGA ALMACÉN").alignment = center_align
        ws.cell(row=y_f+1, column=4, value="RECIBE (FIRMA)").alignment = center_align # <-- Cambio a col 4

        writer.close()

    def mostrar_vista_previa_lote(self, cursor, municipios, per, anio):
        """Actualiza la tabla con el resumen del lote antes de procesar"""
        for i in self.tabla.get_children(): self.tabla.delete(i)
        
        # Ajustamos los encabezados para las 4 columnas
        self.tabla.heading("SKU", text="MUNICIPIO")
        self.tabla.heading("TIPO", text=" ") # Dejamos el encabezado de TIPO vacío para el resumen
        self.tabla.heading("MEDICAMENTO", text="ESTADO DE CARGA")
        self.tabla.heading("CANTIDAD", text="TOTAL PRODUCTOS")

        for mun in municipios:
            cursor.execute("SELECT COUNT(*) FROM historial_retornos WHERE municipio=? AND periodo=? AND anio=?", (mun, per, anio))
            total = cursor.fetchone()[0]
            
            # Insertamos 4 valores, dejando el segundo en blanco (correspondiente a TIPO)
            self.tabla.insert("", "end", values=(mun, "", "LISTO PARA PROCESAR", f"{total} SKUs"))
    
    # Reemplaza la función completa
    def reset_tabla_original(self):
        self.tabla.heading("SKU", text="SKU")
        self.tabla.heading("TIPO", text="TIPO")
        self.tabla.heading("MEDICAMENTO", text="MEDICAMENTO")
        self.tabla.heading("CANTIDAD", text="CANTIDAD")
    def regresar(self):
        self.root.destroy()
        self.main_window.deiconify()